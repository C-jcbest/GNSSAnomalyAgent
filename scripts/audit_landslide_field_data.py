"""Inspect the public Oak Ridge workbook without inventing ENU or unit conversions."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


def audit(workbook_path, out):
    out.mkdir(parents=True, exist_ok=False)
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheets = {}
    for sheet in workbook:
        rows = list(sheet.iter_rows(values_only=True))
        headers, data = rows[0], rows[1:]
        export_name = sheet.title.replace(" ", "-") + "-raw.csv"
        with (out / export_name).open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(headers)
            writer.writerows(data)
        decimal_column = 1 if sheet.title in {"OREO", "ORE2"} else 0
        decimals = [row[decimal_column] for row in data
                    if isinstance(row[decimal_column], (int, float))]
        values = [row[-1] for row in data if isinstance(row[-1], (int, float))]
        result = {"headers": list(headers), "data_rows": len(data),
                  "decimal_year_range": [min(decimals), max(decimals)],
                  "decimal_year_nonincreasing_pairs": sum(b <= a for a, b in
                                                          zip(decimals, decimals[1:])),
                  "numeric_measurement_rows": len(values),
                  "measurement_range_source_units": [min(values), max(values)],
                  "missing_cells_by_column": [sum(row[index] is None for row in data)
                                              for index in range(len(headers))],
                  "first_rows": [[str(value) if isinstance(value, datetime) else value
                                  for value in row] for row in data[:4]],
                  "export": export_name}
        if sheet.title in {"OREO", "ORE2"}:
            dates = [row[0] for row in data if isinstance(row[0], datetime)]
            result["date_column_nonincreasing_pairs"] = sum(b <= a for a, b in
                                                            zip(dates, dates[1:]))
            result["date_column_year_decimal_year_disagreement_rows"] = sum(
                isinstance(row[0], datetime) and isinstance(row[1], (int, float))
                and row[0].year != int(row[1]) for row in data)
            result["gps_scope"] = "Scalar corrected 3D displacement, not separate N/E/U"
            result["gps_unit_status"] = "Not stated in inspected workbook header; not converted"
        sheets[sheet.title] = result
    workbook.close()
    report = {"source_sha256": hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
              "license": "CC-BY-4.0 per Zenodo metadata",
              "site_count": 1, "gps_station_count": 2,
              "sheets": sheets,
              "decision": "Preserve raw values. Resolve date encoding and GPS units from "
                          "authoritative source before preparing a common-task field input. "
                          "No independent activity/stage labels are supplied."}
    (out / "audit.json").write_text(json.dumps(report, ensure_ascii=False, indent=2),
                                   encoding="utf-8")
    lines = ["# Oak Ridge 数据可用性核查", "", "来源：Zenodo 11511667，CC-BY-4.0。", "",
             "同一滑坡两个GPS站，不能视为两个独立场地。原文件只给校正后的三维位移标量，"
             "不能伪造为N/E/U三分量；没有独立活动或阶段参考。", "",
             "| 工作表 | 数据行 | 十进制年份范围 | 数值观测行 |",
             "|---|---:|---|---:|"]
    for name, row in sheets.items():
        lines.append(f"| {name} | {row['data_rows']} | {row['decimal_year_range']} | "
                     f"{row['numeric_measurement_rows']} |")
    lines.extend(["", "GPS日期列与十进制年份存在明显冲突；不得直接使用Excel日期列开展日历评分。"
                  "GPS位移单位在所检表头未注明，暂未换算为毫米。保留原CSV与书目，"
                  "下一步需核查原文或作者数据说明，确认日期编码、单位和缺测后再适配。",
                  "", "本核查不产生现场识别成绩，也不修改源工作簿。"])
    (out / "audit.md").write_text("\n".join(lines), encoding="utf-8")
    print(json.dumps(sheets, ensure_ascii=True, indent=2))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    audit(args.workbook, args.out)


if __name__ == "__main__":
    main()
