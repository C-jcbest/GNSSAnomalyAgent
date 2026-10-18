"""Build the evidence-updated paper using the existing manuscript's Word typesetting helpers."""

from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

import build_manuscript as typeset
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from lxml import etree

ROOT = Path(__file__).resolve().parent
REVISION = ROOT / "revisions/range-20261018"
OUTPUT = ROOT / "GNSS多类型异常识别论文初稿_范围视觉与形变更新版.docx"
TITLE = "面向GNSS多类型异常识别的视觉数值分工方法"
FIGURES = {
    "framework": "视觉暂态范围与数值点类及形变的职责分配",
    "metrics": "暂态范围主指标与全站辅助日指标",
    "modality": "无数值回填时图像与文字输出的正确检出和误报",
    "strata": "暂态阳性场景表现及无暂态场景误报",
    "positive": "e0135可观测暂态补检与缺测标签",
    "negative": "e0260永久阶跃平台误报为暂态的反例",
    "deformation": "108站开发确认的形变性能保留与边界消融",
}
COMMENTS = {
    "3 视觉数值分工方法": "本章替代旧SDR-VLM语义修复主线，与EXP37实际流程对应；局部向量配对属于既有数值组件，未改名为视觉创新。",
    "4.2 分类型评价与不确定性": "按当前评价政策分别统计四类。已删除旧稿四类等权F1主指标，范围Aff均值与阴性误报需并读。",
    "4.7 形变性能保留补充实验": "本节对应EXP45的108个新开发确认站，不是360站独立测试的一部分；最终视觉修边0站，持平成绩归属数值。",
    "6 结论": "结论限于仿真条件下的暂态范围收益及形变性能保留，未使用SOTA或实测预警有效等无证据表述。",
}


def main():
    chapters = [
        (p, p.read_text(encoding="utf-8")) for p in sorted((REVISION / "chapters").glob("*.md"))
    ]
    entire = "\n\n".join(text for _, text in chapters)
    citation_order = list(dict.fromkeys(re.findall(r"@([A-Za-z0-9]+)", entire)))
    citations = {key: i + 1 for i, key in enumerate(citation_order)}
    references = {
        r["id"]: r
        for r in json.loads(
            (REVISION / "working/extra-references.json").read_text(encoding="utf-8")
        )["references"]
    }
    assert set(citations) <= references.keys()
    tables = {
        key: i + 1 for i, key in enumerate(dict.fromkeys(re.findall(r"^表\[(\w+)\]", entire, re.M)))
    }
    table_data = json.loads((REVISION / "working/tables.json").read_text(encoding="utf-8"))
    transform = etree.XSLT(
        etree.parse("C:/Program Files/Microsoft Office/root/Office16/MML2OMML.XSL")
    )
    typeset.MATH_TRANSFORM = transform
    doc = Document()
    typeset.set_style(doc)
    doc.core_properties.title = TITLE
    doc.core_properties.subject = "GNSS范围视觉识别与形变性能保留"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "GNSS;OWV-GNSS;Affiliation;视觉数值分工"
    equation = 0
    for path, chapter in chapters:
        lines = chapter.splitlines()
        index = 0
        while index < len(lines):
            line = lines[index].strip()
            index += 1
            if not line:
                continue
            if line == "$$":
                formula = []
                while lines[index].strip() != "$$":
                    formula.append(lines[index])
                    index += 1
                index += 1
                equation += 1
                typeset.add_equation(doc, " ".join(formula), equation, transform)
                continue
            token = re.fullmatch(r"\[\[(FIGURE|TABLE):(\w+)\]\]", line)
            if token:
                kind, key = token.groups()
                if kind == "TABLE":
                    typeset.add_table(doc, table_data[key], citations, tables)
                    table = doc.tables[-1]
                    headers = table_data[key][0]
                    if "TP FP FN日" in headers:
                        count_column = headers.index("TP FP FN日")
                        widths = []
                        for column_index in range(len(headers)):
                            if column_index == 0:
                                widths.append(3.7)
                            elif column_index == count_column:
                                widths.append(4.1)
                            else:
                                widths.append(8.2 / (len(headers) - 2))
                        for column, width in zip(table.columns, widths):
                            column.width = Cm(width)
                            for cell in column.cells:
                                cell.width = Cm(width)
                else:
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.first_line_indent = Pt(0)
                    # Exact body line spacing would clip an inline picture to one text line.
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.keep_with_next = True
                    p.add_run().add_picture(
                        str(REVISION / "figures" / f"{key}.png"), width=Cm(15.6)
                    )
                    number = list(FIGURES).index(key) + 1
                    p = doc.add_paragraph(f"图{number}  {FIGURES[key]}", "Caption")
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                continue
            if line.startswith("|"):
                rows = []
                while True:
                    values = [v.strip() for v in line.strip("|").split("|")]
                    if not all(re.fullmatch(r"[-: ]+", v) for v in values):
                        rows.append(values)
                    if index >= len(lines) or not lines[index].strip().startswith("|"):
                        break
                    line = lines[index].strip()
                    index += 1
                typeset.add_table(doc, rows, citations, tables)
                continue
            if line.startswith("#"):
                hashes, content = line.split(" ", 1)
                if path.name.startswith("00") and len(hashes) == 1:
                    style = "Title"
                elif len(hashes) == 1:
                    style = "Heading 1"
                else:
                    style = "Heading 2"
                p = doc.add_paragraph(content, style)
                if path.name.startswith("00"):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if content.startswith("Visual and Numerical") or content == "1 引言":
                    p.paragraph_format.page_break_before = True
                if content in COMMENTS:
                    doc.add_comment(p.runs, text=COMMENTS[content], author="修订说明", initials="R")
                continue
            is_caption = bool(re.match(r"^表\[\w+\]", line))
            p = doc.add_paragraph(style="Caption" if is_caption else "Normal")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_caption else WD_ALIGN_PARAGRAPH.JUSTIFY
            if is_caption:
                p.paragraph_format.keep_with_next = True
            if line.startswith(("关键词", "Keywords:")):
                p.paragraph_format.first_line_indent = Pt(0)
            typeset.text_with_markup(p, line, citations, tables)
    doc.add_paragraph("参考文献", "Heading 1")
    for key in citation_order:
        p = doc.add_paragraph(
            typeset.reference_text(references[key], citations[key]), "Bibliography"
        )
        p.paragraph_format.left_indent = Cm(0.7)
        p.paragraph_format.first_line_indent = Cm(-0.7)
        p.paragraph_format.keep_together = True
    doc.save(OUTPUT)
    manifest = {
        "docx": str(OUTPUT),
        "sha256": hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),
        "chapters": {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p, _ in chapters},
        "figures": FIGURES,
        "tables": tables,
        "equations": equation,
        "citations": citations,
        "scope": "EXP37 independent test and separate EXP45 development confirmation; no new inference",
        "characters": len(entire),
        "body_paragraphs": len(doc.paragraphs),
        "word_comments": len(COMMENTS),
    }
    (REVISION / "working/build-manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
