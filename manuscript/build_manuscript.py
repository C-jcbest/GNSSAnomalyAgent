"""Build the Chinese paper from reviewed chapters using the bundled Python runtime.

Equations use latex2mathml 3.81.1 in a task-local temporary dependency directory
and Microsoft's installed MathML-to-OMML stylesheet. No experiment is rerun.
"""

from pathlib import Path
import json
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parent
ASSETS = ROOT / "figures/exp06"
OUTPUT = ROOT / "GNSS多类型异常识别论文初稿_SDR-VLM实验更新版.docx"
sys.path.insert(0, str(Path(os.environ["TEMP"]) / "gnss-paper-authoring-deps"))
from latex2mathml.converter import convert as latex_to_mathml


from prepare_sdr_figures import prepare_figures


def font_properties(font, chinese="宋体", size=12, bold=False):
    font.name = "Times New Roman"
    font.size = Pt(size)
    font.bold = bold
    font.color.rgb = RGBColor(0, 0, 0)
    fonts = font._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:eastAsia"), chinese)
    for attribute in ["asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"]:
        fonts.attrib.pop(qn(f"w:{attribute}"), None)


def set_style(document):
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2.3), Cm(2.3)
    section.left_margin, section.right_margin = Cm(2.5), Cm(2.5)
    section.header_distance, section.footer_distance = Cm(1.0), Cm(1.1)
    normal = document.styles["Normal"]
    font_properties(normal.font)
    normal.paragraph_format.line_spacing = Pt(18.5)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.first_line_indent = Pt(24)
    normal.paragraph_format.widow_control = True
    for name, size in [("Title", 19), ("Heading 1", 15), ("Heading 2", 12.5)]:
        style = document.styles[name]
        font_properties(style.font, "黑体", size, bold=True)
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.25
        style.paragraph_format.keep_with_next = True
    for name, size in [("Caption", 10.5), ("Bibliography", 10.5)]:
        if name not in document.styles:
            document.styles.add_style(name, 1)
        style = document.styles[name]
        font_properties(style.font, "宋体", size)
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.space_before = Pt(5)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.line_spacing = 1.15
        if name == "Bibliography":
            style.paragraph_format.space_before = Pt(0)
            style.paragraph_format.space_after = Pt(3)
            style.paragraph_format.line_spacing = 1.1
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Pt(0)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    header = section.header.paragraphs[0]
    header.text = "GNSS多类型异常识别"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.first_line_indent = Pt(0)
    font_properties(header.runs[0].font, size=9)
    for style in document.styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)


FIGURES = {
    "framework": (1, "SDR-VLM的独立数值路径与受限视觉语义修复", 16),
    "comparison_overall": (2, "48个测试站上主要方法的四类等权日F1与事件F1", 15.6),
    "comparison_types": (3, "48个测试站上主要方法的分类型日F1", 15.6),
    "case_transient": (4, "首个视觉修复案例的原始曲线与阶跃和暂态标签对照", 15.6),
    "case_deformation": (5, "首个纯形变站的原始位移与活动边界对照", 15.6),
}


def text_with_markup(paragraph, text, citation_numbers, table_numbers):
    text = re.sub(r"表\[\s*(\w+)\s*\]", lambda m: f"表{table_numbers[m[1]]}", text)
    text = text.replace("[@Andrade2024]", "[@Andrade2024preprint]")
    pattern = r"(\[@[^\]]+\]|\*\*[^*]+\*\*|`[^`]+`)"
    for token in re.split(pattern, text):
        if not token:
            continue
        if token.startswith("[@"):
            keys = re.findall(r"@([A-Za-z0-9]+)", token)
            run = paragraph.add_run("[" + ",".join(str(citation_numbers[k]) for k in keys) + "]")
            run.font.superscript = True
            run.font.size = Pt(9)
        elif token.startswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("`"):
            paragraph.add_run(token[1:-1])
        else:
            add_inline_math(paragraph, token)


def add_inline_math(paragraph, text):
    """Typeset the author's inline symbols without changing implementation names."""
    pattern = r"(?<![A-Za-z_])(?:TP|FP|FN|RSS|x̂|[xarcempqvbutngNEUVRFQLHWOPDKAℓλεσκτπŷ])_(?:\{[^}]+\}|[A-Za-z0-9τ]+)(?:\^(?:\{[^}]+\}|[A-Za-z0-9]+))?"
    cursor = 0
    for match in re.finditer(pattern, text):
        paragraph.add_run(text[cursor:match.start()])
        expression = match[0].replace("x̂", r"\widehat{x}").replace("ŷ", r"\widehat{y}").replace("ℓ", r"\ell").replace("σ", r"\sigma").replace("π", r"\pi").replace("λ", r"\lambda")
        expression = expression.replace("ε", r"\varepsilon").replace("κ", r"\kappa").replace("τ", r"\tau")
        expression = re.sub(r"_([A-Za-z0-9]+)", lambda item: "_{" + (r"\mathrm{" + item[1] + "}" if len(item[1]) > 1 else item[1]) + "}", expression)
        expression = re.sub(r"\^([A-Za-z0-9]+)", lambda item: "^{" + (r"\mathrm{" + item[1] + "}" if len(item[1]) > 1 else item[1]) + "}", expression)
        mathml = etree.fromstring(latex_to_mathml(expression).encode())
        paragraph._p.append(MATH_TRANSFORM(mathml).getroot())
        cursor = match.end()
    paragraph.add_run(text[cursor:])


def add_table(document, rows, citations, tables):
    count = len(rows[0])
    table = document.add_table(rows=1, cols=count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    # Narrative columns need more room; metric columns remain compact and aligned.
    if count == 2:
        widths = [3.7, 12.3]
    elif count == 3 and "bootstrap" in rows[0][-1]:
        widths = [5.2, 4.3, 6.5]
    elif count == 3:
        widths = [9, 3.5, 3.5]
    elif count == 5 and "批次" in rows[0][0]:
        widths = [4.0, 4.0, 2.5, 3.0, 2.5]
    elif count == 4 and "异常类型" in rows[0][0]:
        widths = [2, 4.6, 4.9, 4.5]
    else:
        first = 4.9 if count >= 6 else 5.0
        widths = [first] + [(16 - first) / (count - 1)] * (count - 1)
    borders = OxmlElement("w:tblBorders")
    for side in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        border = OxmlElement(f"w:{side}")
        for attr, value in [("val", "single"), ("sz", "4"), ("color", "D9D9D9")]:
            border.set(qn(f"w:{attr}"), value)
        borders.append(border)
    table._tbl.tblPr.append(borders)
    for index, values in enumerate(rows):
        row = table.rows[0] if index == 0 else table.add_row()
        row_properties = row._tr.get_or_add_trPr()
        row_properties.append(OxmlElement("w:cantSplit"))
        if index == 0:
            row_properties.append(OxmlElement("w:tblHeader"))
        for col, value in enumerate(values):
            cell = row.cells[col]
            cell.width = Cm(widths[col])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side in ["top", "bottom", "left", "right"]:
                item = OxmlElement(f"w:{side}")
                item.set(qn("w:w"), "85")
                item.set(qn("w:type"), "dxa")
                margins.append(item)
            tcpr.append(margins)
            if index == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "EDEFF1")
                tcpr.append(shade)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.1
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.keep_with_next = index == 0 or (len(rows) <= 5 and index < len(rows) - 1)
            numeric = bool(re.fullmatch(r"[0-9.±—–− /%]+", value))
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 or numeric else WD_ALIGN_PARAGRAPH.LEFT
            text_with_markup(paragraph, value, citations, tables)
            for run in paragraph.runs:
                font_properties(run.font, size=10.5, bold=index == 0 or run.bold)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(3)
    spacer.paragraph_format.first_line_indent = Pt(0)


def add_equation(document, latex, number, transform):
    latex = latex.replace(r"\arg\max", r"\mathop{\mathrm{arg\,max}}")
    latex = latex.replace(r"\arg\min", r"\mathop{\mathrm{arg\,min}}")
    # A multi-expression display is laid out over two lines when needed.
    pieces = [latex]
    if len(latex) > 180:
        parts = re.split(r",\\(?:qquad|quad)\s*", latex)
        if len(parts) > 1:
            midpoint = (len(parts) + 1) // 2
            pieces = [r",\quad ".join(parts[:midpoint]), r",\quad ".join(parts[midpoint:])]
    for index, expression in enumerate(pieces):
        p = document.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.1
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = index < len(pieces) - 1
        mathml = etree.fromstring(latex_to_mathml(expression.strip()).encode())
        result = transform(mathml).getroot()
        p._p.append(result)
        if index == len(pieces) - 1:
            p.add_run(f"   （{number}）")


def reference_text(reference, number):
    authors = reference["authors"]
    author_text = ", ".join(authors[:3])
    if len(authors) > 3:
        author_text += ", et al."
    source = reference["source"]
    details = str(reference["year"])
    if reference.get("volume"):
        details += ", " + reference["volume"]
        if reference.get("issue"):
            details += f"({reference['issue']})"
    if reference.get("pages") or reference.get("article_number"):
        details += ": " + reference.get("pages", reference.get("article_number", ""))
    if reference.get("arxiv_id") and "arXiv" in source:
        details += ", arXiv:" + reference["arxiv_id"]
    title = reference["title"]
    title_end = "" if title.endswith(("?", ".", "!")) else "."
    text = f"[{number}] {author_text.rstrip('.')}. {title}{title_end} {source}, {details}."
    if reference.get("doi"):
        text += " DOI: " + reference["doi"] + "."
    elif "arXiv" not in source:
        text += " " + reference["url"]
    return text


def main():
    global MATH_TRANSFORM
    prepare_figures()
    paths = sorted((ROOT / "chapters").glob("*.md"))
    chapters = [(path, path.read_text(encoding="utf-8-sig")) for path in paths]
    entire = "\n\n".join(text for _, text in chapters).replace("[@Andrade2024]", "[@Andrade2024preprint]")
    citation_order = list(dict.fromkeys(re.findall(r"@([A-Za-z0-9]+)", entire)))
    citations = {key: index + 1 for index, key in enumerate(citation_order)}
    reference_data = json.loads((ROOT / "working/extra-references.json").read_text(encoding="utf-8"))
    references = {entry["id"]: entry for entry in reference_data["references"]}
    missing = set(citations) - set(references)
    if missing:
        raise ValueError(f"Missing verified references: {missing}")
    table_order = list(dict.fromkeys(re.findall(r"^表\[\s*(\w+)\s*\]\s+", entire, re.M)))
    tables = {key: index + 1 for index, key in enumerate(table_order)}
    transform = etree.XSLT(etree.parse("C:/Program Files/Microsoft Office/root/Office16/MML2OMML.XSL"))
    MATH_TRANSFORM = transform
    document = Document()
    set_style(document)
    document.core_properties.title = "融合位移结构分解与视觉语义修复的GNSS多类型异常识别方法"
    document.core_properties.subject = "GNSS-VLM-EXP06-v1"
    document.core_properties.author = ""
    document.core_properties.keywords = "GNSS;视觉语言模型;多类型异常识别;SDR-VLM"
    equation = 0
    for path, chapter in chapters:
        lines = chapter.splitlines()
        index = 0
        while index < len(lines):
            line = lines[index].strip()
            index += 1
            if not line:
                continue
            if line == "$$":
                formula = []
                while index < len(lines) and lines[index].strip() != "$$":
                    formula.append(lines[index])
                    index += 1
                index += 1
                equation += 1
                add_equation(document, " ".join(formula), equation, transform)
                continue
            figure = re.fullmatch(r"\[\[FIGURE:(\w+)\]\]", line)
            if figure:
                key = figure[1]
                number, caption, width = FIGURES[key]
                p = document.add_paragraph()
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(ASSETS / f"{key}.png"), width=Cm(width))
                p = document.add_paragraph(f"图{number}  {caption}", "Caption")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                continue
            if line.startswith("|"):
                rows = []
                table_lines = [line]
                while index < len(lines) and lines[index].strip().startswith("|"):
                    table_lines.append(lines[index].strip())
                    index += 1
                for table_line in table_lines:
                    values = [part.strip() for part in table_line.strip("|").split("|")]
                    if all(re.fullmatch(r"[-: ]+", value) for value in values):
                        continue
                    rows.append(values)
                add_table(document, rows, citations, tables)
                continue
            if line.startswith("#"):
                hashes, content = line.split(" ", 1)
                if path.name.startswith("00"):
                    if len(hashes) == 1:
                        style = "Title"
                    else:
                        style = "Heading 2"
                else:
                    style = "Heading 1" if len(hashes) == 1 else "Heading 2"
                if style == "Title":
                    content = content.replace("的GNSS", "的\nGNSS")
                p = document.add_paragraph(content, style)
                if path.name.startswith("00"):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    if content.startswith("GNSS Multi-Type"):
                        p.paragraph_format.space_before = Pt(12)
                        p.paragraph_format.page_break_before = True
                if path.name.startswith("01") and len(hashes) == 1:
                    p.paragraph_format.page_break_before = True
                continue
            is_caption = bool(re.match(r"^表\[\s*\w+\s*\]\s+", line))
            p = document.add_paragraph(style="Caption" if is_caption else "Normal")
            if is_caption:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next = True
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if line.startswith(("关键词", "Keywords:")):
                p.paragraph_format.first_line_indent = Pt(0)
            text_with_markup(p, line, citations, tables)
    document.add_paragraph("参考文献", "Heading 1")
    for key in citation_order:
        p = document.add_paragraph(reference_text(references[key], citations[key]), "Bibliography")
        p.paragraph_format.left_indent = Cm(.7)
        p.paragraph_format.first_line_indent = Cm(-.7)
        p.paragraph_format.keep_together = True
    document.save(OUTPUT)
    manifest = {
        "docx": str(OUTPUT), "chapters": [path.name for path, _ in chapters],
        "citations": citations, "tables": tables, "equations": equation,
        "figures": {key: value[0] for key, value in FIGURES.items()},
        "experiment": "GNSS-VLM-EXP06-v1",
    }
    (ROOT / "working/build-manifest-sdr.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(manifest, ensure_ascii=False))


if __name__ == "__main__":
    main()
