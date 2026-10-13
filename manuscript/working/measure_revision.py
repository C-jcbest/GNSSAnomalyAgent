"""Record the revised manuscript's section allocation and document structure."""

import hashlib
import json
from pathlib import Path
import re
import subprocess
from zipfile import ZipFile

from docx import Document
from lxml import etree
import pdfplumber

from measure_section_allocation import add_percentages, pdf_sections


ROOT = Path(__file__).resolve().parents[1]
WORKING = ROOT / "working"
DOCX = ROOT / "GNSS多类型异常识别论文初稿_结构优化版.docx"
PDF = WORKING / "render-revision/paper.pdf"


def main():
    starts = []
    reference_start = None
    with pdfplumber.open(PDF) as document:
        page_count = len(document.pages)
        for number, page in enumerate(document.pages, 1):
            for line in page.extract_text_lines():
                title = line["text"]
                if re.match(r"^[1-6] (引言|相关工作|GNSS|结果与分析|讨论|结论)", title):
                    starts.append((title, number, 0, line["top"]))
                elif title == "参考文献":
                    reference_start = (number, 0, line["top"])
    assert len(starts) == 6 and reference_start is not None
    sections = pdf_sections(PDF, starts, reference_start, 67, 778)
    add_percentages(sections)

    source_counts = []
    for chapter in sorted((ROOT / "chapters").glob("*.md")):
        if chapter.name.startswith("00"):
            continue
        text = chapter.read_text(encoding="utf-8")
        row = {
            "source": str(chapter.relative_to(ROOT)),
            "chinese_characters": len(re.findall(r"[\u4e00-\u9fff]", text)),
        }
        if chapter.name.startswith(("01", "02")):
            old = subprocess.check_output(
                ["git", "show", f"f37f140:manuscript/chapters/{chapter.name}"],
                cwd=ROOT.parent,
            ).decode("utf-8")
            previous_count = len(re.findall(r"[\u4e00-\u9fff]", old))
            row["previous_chinese_characters"] = previous_count
            row["increase_percent"] = 100 * (row["chinese_characters"] / previous_count - 1)
        source_counts.append(row)

    allocation = {
        "date": "2026-10-13",
        "method": "Alphabetic PDF characters, body only; page spans include figures and layout. Same counting method as the previous allocation study.",
        "pdf_sha256": hashlib.sha256(PDF.read_bytes()).hexdigest(),
        "sections": sections,
        "source_chinese_character_counts": source_counts,
    }
    (WORKING / "section-allocation-revision.json").write_text(
        json.dumps(allocation, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    manuscript = Document(DOCX)
    with ZipFile(DOCX) as archive:
        xml = etree.fromstring(archive.read("word/document.xml"))
    ns = {
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
    }
    all_text = "".join(xml.xpath("//w:t/text() | //m:t/text()", namespaces=ns))
    old_docx = ROOT / "GNSS多类型异常识别论文初稿.docx"
    old_hash = hashlib.sha256(old_docx.read_bytes()).hexdigest()
    manifest = json.loads((WORKING / "build-manifest-revision.json").read_text(encoding="utf-8"))
    integrity = {
        "pages_from_word_render": page_count,
        "headings": [p.text for p in manuscript.paragraphs if p.style.name == "Heading 1"],
        "tables": len(manuscript.tables),
        "figures": len(manuscript.inline_shapes),
        "display_equation_groups": manifest["equations"],
        "references": len(manifest["citations"]),
        "editable_math_objects": len(xml.xpath("//m:oMath", namespaces=ns)),
        "unresolved_placeholders": re.findall(r"\[\[FIGURE[^\]]*|\[@[^\]]+|表\s*\[[^\]]+|TODO|TBD", all_text),
        "docx_size_bytes": DOCX.stat().st_size,
        "docx_sha256": hashlib.sha256(DOCX.read_bytes()).hexdigest(),
        "original_docx_sha256": old_hash,
        "original_docx_unchanged": old_hash == "cea0435704c8f7068097c60579a69bb9525f666073a65d3f2b696f3d3aa465a4",
        "rendered_page_sha256": {
            p.name: hashlib.sha256(p.read_bytes()).hexdigest()
            for p in sorted((WORKING / "render-revision").glob("page-*.png"))
        },
    }
    assert 20 <= page_count <= 30
    assert integrity["tables"] == 8 and integrity["figures"] == 5
    assert integrity["original_docx_unchanged"] and not integrity["unresolved_placeholders"]
    (WORKING / "final-integrity-revision.json").write_text(
        json.dumps(integrity, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    body = "\n\n".join(
        chapter.read_text(encoding="utf-8")
        for chapter in sorted((ROOT / "chapters").glob("*.md"))
        if not chapter.name.startswith("00")
    )
    (WORKING / "draft-revision.md").write_text(body, encoding="utf-8")
    print(json.dumps({"pages": page_count, "sections": sections, "source_counts": source_counts}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
